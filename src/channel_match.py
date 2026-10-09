"""Match channels from channel_lib against input Excel file. Output JSON summary.
同名多渠道时用金额消歧：优先 商品单价，其次 支付金额+平台补贴金额（与渠道名 -99/-49/-199 后缀比对）；
无法判定则留在未匹配列表（页面手动补齐）。规格=默认 的行不需要匹配，直接从成品表中删除。
商品单价填充（同订单池规则，2026-10-07）：单价缺失（空/#N/A/-/无）时用 支付金额+平台补贴金额 计算，
仅认可档位（49/99/119/199/219）才回填；源表没有单价列时在支付金额前新增一列。
"""
import sys, os, sqlite3, json, re
import pandas as pd
from openpyxl.utils import get_column_letter

# 与 web-chat/full_pool.py 的 VALID_UNIT_PRICES / _NA_PLACEHOLDERS 保持一致
VALID_UNIT_PRICES = {49.0, 99.0, 119.0, 199.0, 219.0}
MISSING_TOKENS = {"", "-", "无", "#N/A", "N/A", "NA", "#NA", "#VALUE!", "#REF!", "#NULL!", "NULL", "NONE"}

def price_missing(x):
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return True
    if isinstance(x, (int, float)) and float(x) == 0:
        return True   # 数值 0 同订单池按缺失处理（or "" 语义）
    return str(x).strip().upper() in MISSING_TOKENS

def to_num(x):
    """'71.40元' / 71.4 / '1,234.00元' -> float；解析失败返回 None"""
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return None
    s = str(x).replace("元", "").replace(",", "").strip()
    try:
        return float(s)
    except ValueError:
        return None

def price_of(channel):
    """'初中数学-99' -> 99.0；'高中内容直播' -> None"""
    m = re.search(r"-(\d+(?:\.\d+)?)$", str(channel).strip())
    return float(m.group(1)) if m else None

def main():
    file_path = sys.argv[1]
    db_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "data", "kehu.db")

    # Load channel library
    conn = sqlite3.connect(db_path)
    lib = pd.read_sql("SELECT prod_name, prod_id, channel FROM channel_lib", conn)
    conn.close()

    # Lookup: prod_id / prod_name -> [channel, ...]（同名/同ID 多渠道保留全部，匹配时用金额消歧）
    id_map = {}
    name_map = {}
    for _, row in lib.iterrows():
        pid = str(row['prod_id']).strip()
        pname = str(row['prod_name']).strip()
        ch = row['channel']
        if pid and pid not in id_map:
            id_map[pid] = []
        if pid and ch not in id_map[pid]:
            id_map[pid].append(ch)
        if pname and pname not in name_map:
            name_map[pname] = []
        if pname and ch not in name_map[pname]:
            name_map[pname].append(ch)

    # Read input
    df = pd.read_excel(file_path)
    # Find 订单编号 / 商品名称 / 商品ID / 规格 / 金额 columns
    order_col = None
    name_col = None
    id_col = None
    spec_col = None
    price_col = None
    pay_col = None
    sub_col = None
    for c in df.columns:
        cs = str(c).strip()
        if ('订单编号' in cs or '订单号' in cs or '订单ID' in cs) and order_col is None:
            order_col = c
        if '商品名称' in cs and name_col is None:
            name_col = c
        if '商品ID' in cs and id_col is None:
            id_col = c
        if cs == '规格' and spec_col is None:
            spec_col = c
        if '商品单价' in cs and price_col is None:
            price_col = c
        if '支付金额' in cs and pay_col is None:
            pay_col = c
        if '补贴金额' in cs and sub_col is None:
            sub_col = c

    if name_col is None or id_col is None:
        fn = [str(c) for c in df.columns]
        print(json.dumps({"ok": False, "error": f"找不到商品名称/商品ID列。列名: {fn}"}, ensure_ascii=False))
        return

    # 商品单价填充（同订单池规则）：缺失时用 支付金额+平台补贴金额，仅认可档位才回填
    filled_price = 0
    if pay_col is not None or sub_col is not None:
        if price_col is None:
            anchor = pay_col if pay_col is not None else sub_col
            price_col = "商品单价"
            df.insert(df.columns.get_loc(anchor), price_col, "")
        for i, row in df.iterrows():
            if not price_missing(row[price_col]):
                continue
            p = to_num(row[pay_col]) if pay_col is not None else 0.0
            s = to_num(row[sub_col]) if sub_col is not None else 0.0
            if p is None and s is None:
                continue
            total = round((p or 0.0) + (s or 0.0), 2)
            if total in VALID_UNIT_PRICES:
                df.at[i, price_col] = f"{total:.2f}元"
                filled_price += 1

    # Convert key columns to string to preserve full precision (19-digit IDs / order numbers)
    def to_str(x):
        if pd.isna(x):
            return ''
        if isinstance(x, float):
            return str(int(x))
        return str(x).strip()
    if order_col is not None:
        df[order_col] = df[order_col].apply(to_str)
    df[id_col] = df[id_col].apply(to_str)

    # 剔除不需要匹配的行：规格=默认（如赠品/无渠道商品），直接从成品表中删除
    removed_spec = 0
    if spec_col is not None:
        mask = df[spec_col].apply(lambda x: str(x).strip() == '默认' if pd.notna(x) else False)
        removed_spec = int(mask.sum())
        if removed_spec:
            df = df[~mask].reset_index(drop=True)

    def resolve(pname, pid, price, amount):
        cands = (id_map.get(pid) if pid else None) or (name_map.get(pname) if pname else None) or []
        if not cands:
            return ""
        if len(cands) == 1:
            return cands[0]
        # 同名/同ID 多渠道：按金额与渠道名价格后缀消歧（单价优先——支付+补贴有约 22% 行因优惠偏离）
        for amt in (price, amount):
            if amt is None:
                continue
            by_price = [c for c in cands if price_of(c) is not None and abs(price_of(c) - amt) < 0.01]
            if by_price:
                return by_price[0]
        # 金额对不上带后缀的渠道时，退回无价格后缀的渠道（如 初中内容直播）
        no_price = [c for c in cands if price_of(c) is None]
        if no_price:
            return no_price[0]
        return ""   # 无法判定 → 未匹配，进手动补齐列表

    # Match
    matched = 0
    unmatched_items = {}
    channels = []
    for _, row in df.iterrows():
        pid = str(row[id_col]).strip() if pd.notna(row[id_col]) else ""
        pname = str(row[name_col]).strip() if pd.notna(row[name_col]) else ""
        price = to_num(row[price_col]) if price_col is not None else None
        amount = None
        if pay_col is not None or sub_col is not None:
            p = to_num(row[pay_col]) if pay_col is not None else 0.0
            s = to_num(row[sub_col]) if sub_col is not None else 0.0
            if p is not None or s is not None:
                amount = (p or 0.0) + (s or 0.0)
        ch = resolve(pname, pid, price, amount)
        channels.append(ch)
        if ch:
            matched += 1
        else:
            key = (pname, pid)
            unmatched_items[key] = unmatched_items.get(key, 0) + 1

    # Add channel column after 商品ID
    id_idx = df.columns.get_loc(id_col)
    df.insert(id_idx + 1, "渠道", channels)

    # Save output with text format for order/id columns
    base, ext = os.path.splitext(file_path)
    out_path = base + "_渠道" + ext
    with pd.ExcelWriter(out_path, engine='openpyxl') as writer:
        df.to_excel(writer, index=False)
        ws = writer.sheets['Sheet1']
        text_cols = []
        if order_col is not None:
            text_cols.append(df.columns.get_loc(order_col) + 1)
        text_cols.append(id_idx + 1)
        for col_idx in set(text_cols):
            letter = get_column_letter(col_idx)
            for r in range(2, ws.max_row + 1):
                ws[f'{letter}{r}'].number_format = '@'

    print(json.dumps({
        "ok": True,
        "total": len(df),
        "removedSpec": removed_spec,
        "priceFilled": filled_price,
        "matched": matched,
        "unmatched": len(df) - matched,
        "unmatchedItems": [
            {"name": k[0], "id": k[1], "rows": v}
            for k, v in sorted(unmatched_items.items(), key=lambda kv: -kv[1])
        ],
        "output": out_path
    }, ensure_ascii=False))

if __name__ == "__main__":
    main()
