#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
未分配退款时段占比分析（渠道 × 5 时段 + 24h 内占比 + 订单去重计数）
=====================================================================
用途：处理「已瘦身」的退款底表（由 preprocess.py 产出，5 列），
      统计「未分配（=分配前）退款」的时段分布。

输入：Excel 文件（必须包含以下 5 列）
  - 订单编号
  - 渠道
  - 下单时间
  - 退款时间
  - 是否分配（值: 已分配 / 未分配）

输出：
  1. Excel 新增工作表「退款时段汇总（未分配）」
  2. 单独一张高清 PNG 群发图（同目录）

口径：
  - 未分配退款 = 是否分配 == 未分配 AND 退款时间非空
  - 退款数 = 满足条件的「订单编号」去重计数（nunique），不是行数
  - 退款时段 = (退款时间 - 下单时间) 单位：小时
  - 时段分桶：<1h / 1-6h / 6-12h / 12-24h / >=24h
  - 渠道顺序：直播间（=子渠道汇总，置顶） + 99 -> 199 -> 49
  - 24h 内占比 = (1h内 + 1-6h + 6-12h + 12-24h) / 总数
  - 异常数据过滤：diff < 0（退款早于下单）的记录直接丢弃

使用：
  python3 refund_summary.py <thin_excel_path>

依赖：
  pip install pandas openpyxl pillow
"""

import sys
import os
import re
import warnings
from pathlib import Path

import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from PIL import Image, ImageDraw, ImageFont

warnings.filterwarnings('ignore')


# ============== 可调参数 ==============
PERIODS  = ['1小时内', '1-6小时', '6-12小时', '12-24小时', '大于24小时']
SHEET_NAME = '退款'                       # 默认读取的工作表
SHEET_OUT = '退款时段汇总（未分配）'       # 生成的工作表名

ALLOCATED   = '已分配'
UNALLOCATED = '未分配'

# 渠道展示顺序偏好（保持与原 skill 约定一致：99 -> 199 -> 49）
CHANNEL_ORDER_PREF = ['99', '199', '49']

# PNG 高清缩放因子（2 = 分辨率翻倍）
SCALE = 2

# 颜色
HDR_BLUE = (31, 78, 120)
ACC_RED  = (192, 0, 0)
WHITE    = (255, 255, 255)
GRID     = (153, 153, 153)
BG       = (255, 255, 255)
TITLE_CLR = (0, 0, 0)


# ============== 工具函数 ==============
def get_period(h: float) -> str:
    if h < 1:    return '1小时内'
    if h < 6:    return '1-6小时'
    if h < 12:   return '6-12小时'
    if h < 24:   return '12-24小时'
    return '大于24小时'


def find_chinese_font() -> str:
    candidates = [
        '/System/Library/Fonts/PingFang.ttc',
        '/System/Library/Fonts/STHeiti Medium.ttc',
        '/System/Library/Fonts/STHeiti Light.ttc',
        '/System/Library/Fonts/Hiragino Sans GB.ttc',
        '/Library/Fonts/Songti.ttc',
        '/System/Library/Fonts/Supplemental/Songti.ttc',
        'C:/Windows/Fonts/msyh.ttc',
        'C:/Windows/Fonts/simhei.ttf',
    ]
    for p in candidates:
        if os.path.exists(p):
            return p
    return None


def map_alloc(v) -> str:
    s = str(v).strip()
    if s in ('已分配', '是', 'Y', 'y', 'YES', 'Yes', 'yes', '1', 'TRUE', 'True', 'true'):
        return ALLOCATED
    if s in ('未分配', '否', 'N', 'n', 'NO', 'No', 'no', '0', 'FALSE', 'False', 'false'):
        return UNALLOCATED
    if s in ('', 'nan', 'None', '#N/A', 'N/A', 'NA', '#NA'):
        return UNALLOCATED
    return s


def channel_sort_key(ch: str):
    m = re.search(r'(\d+)\s*$', ch)
    suffix = m.group(1) if m else ''
    if suffix in CHANNEL_ORDER_PREF:
        return (0, CHANNEL_ORDER_PREF.index(suffix))
    return (1, int(suffix) if suffix.isdigit() else 0)


def infer_live_name(sub_channels) -> str:
    if not sub_channels:
        return '直播间'
    stems = [re.sub(r'[-_]\d+$', '', c) for c in sub_channels]
    if len(set(stems)) == 1 and stems[0]:
        return stems[0] + '直播间'
    return '直播间'


def load_data(path: str) -> pd.DataFrame:
    df = pd.read_excel(path, sheet_name=SHEET_NAME)
    needed = {'订单编号', '渠道', '下单时间', '退款时间', '是否分配'}
    missing = needed - set(df.columns)
    if missing:
        raise ValueError(f'底表缺少必要列: {missing}\n实际列名: {list(df.columns)}')

    df = df[df['退款时间'].notna()].copy()
    df['下单时间_dt'] = pd.to_datetime(df['下单时间'].astype(str).str.strip(), errors='coerce')
    df['退款时间_dt'] = pd.to_datetime(df['退款时间'].astype(str).str.strip(), errors='coerce')
    df = df.dropna(subset=['下单时间_dt', '退款时间_dt'])
    df['diff_hours'] = (df['退款时间_dt'] - df['下单时间_dt']).dt.total_seconds() / 3600
    df = df[df['diff_hours'] >= 0].copy()
    df['退款时段'] = df['diff_hours'].apply(get_period)
    df['是否分配_规范'] = df['是否分配'].apply(map_alloc)
    return df


def compute_crosstab(df: pd.DataFrame, channels: list) -> pd.DataFrame:
    """未分配退款 → 渠道×时段 订单去重计数表（含直播间汇总行）。"""
    sub = df[df['是否分配_规范'] == UNALLOCATED]
    if sub.empty:
        ct = pd.DataFrame(0, index=channels, columns=PERIODS)
    else:
        ct = (sub.groupby(['渠道', '退款时段'])['订单编号']
                 .nunique()
                 .unstack(fill_value=0))
        ct = ct.reindex(columns=PERIODS, fill_value=0)
        ct = ct.reindex(channels, fill_value=0)

    live = channels[0]
    subs = channels[1:]
    ct.loc[live] = ct.loc[subs].sum(axis=0)
    return ct


def write_excel(path: str, ct: pd.DataFrame, channels: list) -> None:
    wb = openpyxl.load_workbook(path)
    if SHEET_OUT in wb.sheetnames:
        del wb[SHEET_OUT]
    ws = wb.create_sheet(SHEET_OUT, 1)

    header_fill = PatternFill('solid', fgColor='1F4E78')
    acc_fill    = PatternFill('solid', fgColor='C00000')
    white_bold  = Font(color='FFFFFF', bold=True, size=11)
    center      = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin        = Side(border_style='thin', color='999999')
    border      = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ['渠道', '退款时段', '未分配退款数（去重订单）', '退款数分时段占比', '24小时以内退款数占比']
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.fill = header_fill
        c.font = white_bold
        c.alignment = center
        c.border = border

    ct = ct.copy()
    ct['__24h'] = ct[PERIODS[:4]].sum(axis=1)
    ct['__total'] = ct[PERIODS].sum(axis=1)
    ct['__24h_pct'] = (ct['__24h'] / ct['__total'].replace(0, pd.NA) * 100).fillna(0).round(0)

    row_map = {}
    start = 2
    for ch in channels:
        row_map[ch] = (start, start + 4)
        start += 5

    for ch in channels:
        s, e = row_map[ch]
        total = int(ct.loc[ch, '__total'])
        p24 = ct.loc[ch, '__24h_pct']
        within_24h_pct = int(p24) if pd.notna(p24) else 0

        ws.merge_cells(start_row=s, start_column=1, end_row=e, end_column=1)
        a = ws.cell(row=s, column=1, value=ch)
        a.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for r in range(s, e + 1):
            ws.cell(row=r, column=1).border = border

        ws.merge_cells(start_row=s, start_column=5, end_row=e, end_column=5)
        ec = ws.cell(row=s, column=5, value=f'{within_24h_pct}%')
        ec.alignment = center
        ec.fill = acc_fill
        ec.font = white_bold
        for r in range(s, e + 1):
            ws.cell(row=r, column=5).border = border
            ws.cell(row=r, column=5).fill = acc_fill

        for i, p in enumerate(PERIODS):
            r = s + i
            cnt = int(ct.loc[ch, p])
            pct = round(cnt / total * 100) if total else 0
            ws.cell(row=r, column=2, value=p).alignment = center
            ws.cell(row=r, column=3, value=cnt).alignment = center
            ws.cell(row=r, column=4, value=f'{pct}%').alignment = center
            for col in (2, 3, 4):
                ws.cell(row=r, column=col).border = border

    ws.column_dimensions['A'].width = 18
    ws.column_dimensions['B'].width = 14
    ws.column_dimensions['C'].width = 22
    ws.column_dimensions['D'].width = 18
    ws.column_dimensions['E'].width = 20
    ws.row_dimensions[1].height = 30
    for r in range(2, 2 + len(channels) * 5):
        ws.row_dimensions[r].height = 22
    ws.freeze_panes = 'A2'
    wb.save(path)


def draw_png(ct: pd.DataFrame, channels: list, out_path: str) -> None:
    font_path = find_chinese_font()
    if not font_path:
        raise RuntimeError('未找到中文字体，无法生成图片')

    S = SCALE
    font_title = ImageFont.truetype(font_path, S * 18)
    font_h = ImageFont.truetype(font_path, S * 15)
    font_b = ImageFont.truetype(font_path, S * 13)

    col_x = [S * v for v in [20, 160, 280, 400, 550]]
    col_w = [S * v for v in [140, 120, 120, 150, 170]]
    row_h = S * 40
    title_h = S * 46

    W = S * 720
    n_rows = 1 + len(channels) * 5
    table_h = n_rows * row_h
    H = title_h + table_h + S * 20

    img = Image.new('RGB', (W, H), BG)
    draw = ImageDraw.Draw(img)

    # 标题
    draw.text((S * 20, 0), '未分配退款时段占比（分配前）', fill=TITLE_CLR, font=font_title)
    y = title_h

    headers = ['渠道', '退款时段', '退款数', '退款数分时段占比', '24小时以内退款数占比']
    for i, txt in enumerate(headers):
        x0c, x1c = col_x[i], col_x[i] + col_w[i]
        draw.rectangle([x0c, y, x1c, y + row_h], fill=HDR_BLUE, outline=GRID)
        bbox = draw.textbbox((0, 0), txt, font=font_h)
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        draw.text((x0c + (col_w[i] - tw) // 2, y + (row_h - th) // 2 - 2), txt, fill=WHITE, font=font_h)
    y += row_h

    ct = ct.copy()
    ct['__24h'] = ct[PERIODS[:4]].sum(axis=1)
    ct['__total'] = ct[PERIODS].sum(axis=1)

    for ch in channels:
        segs = [int(ct.loc[ch, p]) for p in PERIODS]
        total = sum(segs)
        pcts = [round(s / total * 100) for s in segs] if total else [0] * len(segs)
        within_24h = sum(segs[:4])
        within_24h_pct = round(within_24h / total * 100) if total else 0
        h24_str = f'{within_24h_pct}%'

        ch_x0, ch_x1 = col_x[0], col_x[0] + col_w[0]
        draw.rectangle([ch_x0, y, ch_x1, y + 5 * row_h], outline=GRID)
        bbox = draw.textbbox((0, 0), ch, font=font_b)
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        draw.text((ch_x0 + (col_w[0] - tw) // 2, y + (5 * row_h - th) // 2 - 2), ch, fill=(0, 0, 0), font=font_b)

        h24_x0, h24_x1 = col_x[4], col_x[4] + col_w[4]
        draw.rectangle([h24_x0, y, h24_x1, y + 5 * row_h], fill=ACC_RED, outline=GRID)
        bbox = draw.textbbox((0, 0), h24_str, font=font_b)
        tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
        draw.text((h24_x0 + (col_w[4] - tw) // 2, y + (5 * row_h - th) // 2 - 2), h24_str, fill=WHITE, font=font_b)

        for i, p in enumerate(PERIODS):
            ry = y + i * row_h
            for j, val in enumerate([p, str(segs[i]), f'{pcts[i]}%']):
                x0c, x1c = col_x[j + 1], col_x[j + 1] + col_w[j + 1]
                draw.rectangle([x0c, ry, x1c, ry + row_h], outline=GRID)
                bbox = draw.textbbox((0, 0), val, font=font_b)
                tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
                draw.text((x0c + (col_w[j + 1] - tw) // 2, ry + (row_h - th) // 2 - 2), val, fill=(0, 0, 0), font=font_b)

        y += 5 * row_h

    img.save(out_path)


# ============== 主流程 ==============
def main():
    if len(sys.argv) < 2:
        print('用法: python3 refund_summary.py <thin_excel_path>')
        sys.exit(1)

    path = sys.argv[1]
    if not os.path.exists(path):
        print(f'文件不存在: {path}')
        sys.exit(1)

    print(f'[1/3] 读取底表: {path}')
    df = load_data(path)
    print(f'  → 有退款时间的记录: {len(df)} 行；未分配: {(df["是否分配_规范"] == UNALLOCATED).sum()} 行')

    sub_channels = sorted(set(df['渠道'].dropna().unique()), key=channel_sort_key)
    live_name = infer_live_name(sub_channels)
    channels = [live_name] + sub_channels
    print(f'  → 渠道顺序: {channels}')

    print('[2/3] 计算未分配退款订单去重交叉表')
    ct = compute_crosstab(df, channels)
    print(ct.to_string())

    print(f'[3/3] 写入 Excel 工作表 + 高清 PNG')
    write_excel(path, ct, channels)
    png_path = str(Path(path).with_name(Path(path).stem + '_未分配时段占比图.png'))
    draw_png(ct, channels, png_path)
    print(f'  ✓ Excel: {SHEET_OUT}')
    print(f'  ✓ PNG: {png_path}')

    print('\n=== 未分配退款摘要 ===')
    for ch in channels:
        total = int(ct.loc[ch, PERIODS].sum())
        within_24h = int(ct.loc[ch, PERIODS[:4]].sum())
        pct = round(within_24h / total * 100) if total else 0
        print(f'  {ch}：总 {total}，24h内 {within_24h}（{pct}%）')


if __name__ == '__main__':
    main()
